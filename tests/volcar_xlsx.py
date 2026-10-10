# -*- coding: utf-8 -*-
"""Lee un .xlsx con openpyxl y da sus celdas (volcar_celdas): asi la paridad
del Excel (tests/paridad_excel_paso_62.js) compara el de la app con sus
cifras de referencia, leido como lo leeria otro programa.
"""
import datetime

from openpyxl import load_workbook


def volcar_celdas(ruta):
    """Las hojas de un .xlsx y, de cada una, sus celdas: {ref: [tipo,
    valor, formato]}, con tipo «n» (numero), «d» (fecha, en ISO), «b»
    (si/no) o «t» (texto)."""
    wb = load_workbook(ruta, read_only=True)   # solo leer: mas del doble de rapido
    hojas = {}
    for ws in wb.worksheets:
        celdas = {}
        for fila in ws.iter_rows():
            for c in fila:
                v = getattr(c, "value", None)   # las celdas vacias no tienen
                if v is None or v == "":
                    continue
                if isinstance(v, str):
                    celdas[c.coordinate] = ["t", v, None]
                elif isinstance(v, (datetime.datetime, datetime.date)):
                    celdas[c.coordinate] = ["d", v.isoformat(), c.number_format]
                else:
                    celdas[c.coordinate] = ["b" if isinstance(v, bool) else "n", v, c.number_format]
        hojas[ws.title] = celdas
    wb.close()
    return {"hojas": hojas, "orden": list(hojas)}
